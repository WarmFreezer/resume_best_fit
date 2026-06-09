import os
import json
import anthropic
from docx import Document
from datetime import datetime
from dotenv import load_dotenv
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

load_dotenv()

client = anthropic.Anthropic(api_key=os.getenv("ANTHROPIC_API_KEY"))

message = "What is 2 + 2?"

resume = json.load(open("data.json", "r", encoding="utf-8"))
job_description = input("Enter a job description: ")
prompt = "Score each entry in the entries section of the resume json based on the included job description. \
    Consider key words and skills. (0.0-1.0) corresponding to each entry. Return only a dictionary of entries \
    and their scores. Resume as json: " + json.dumps(resume) + ". The job description is: " + job_description

entries = client.messages.create(
    model="claude-haiku-4-5-20251001",
    max_tokens=2048,
    messages=[
        {"role": "user", "content": prompt}
    ]
)

skills = client.messages.create(
    model="claude-haiku-4-5-20251001",
    max_tokens=2048,
    messages=[
        {"role": "user", "content": "Score (0.0-1.0) the skills in this resume: " + json.dumps(resume) + "\
            Against the job description: " + job_description + "Return a json dictionary of the most relevant skills \
            for this job description based on the resume with their scores. Only return the skill and the correlating score.\
            Do not return any other text."}
    ]
)

def extract_json(text):
    text = text.strip()
    if text.startswith("```"):
        text = text.split("```", 2)[1]
        if text.startswith("json"):
            text = text[4:]
        text = text.rsplit("```", 1)[0]
    return text.strip()

entries_dict = json.loads(extract_json(entries.content[0].text))
skills_dict = json.loads(extract_json(skills.content[0].text))

name = resume["name"]
title = resume["title"]
location = resume["location"]
email = resume["email"]
education = resume["education"]

entries = resume["entries"]
# Where the entries have the section "experience"
experience = [entry for entry in entries if entry["section"] == "experience"]
# Where the entries have the section "projects"
projects = [entry for entry in entries if entry["section"] == "projects"]
# Where the entries have the section "leadership"
leadership = [entry for entry in entries if entry["section"] == "leadership"]
# Where the entries have the section "awards"
awards = [entry for entry in entries if entry["section"] == "awards"]
skills_list = resume["skills"]

average_score = sum(entries_dict.values()) / len(entries_dict) if entries_dict else 0
print(f"Average Score: {average_score}")

if average_score > 0.5:
    relevant_experience = []
    for experience_entry in experience:
        if experience_entry["title"] in entries_dict:
            score = entries_dict[experience_entry["title"]]
            if score > 0.5:
                relevant_experience.append(experience_entry)
                print(f"Experience: {experience_entry['title']}, Score: {score}")

    if len(relevant_experience) == 0:
        # Append the highest scoring experience entry if no entries are above the threshold
        highest_scoring_entry = max(experience, key=lambda x: entries_dict.get(x["title"], 0))
        relevant_experience.append(highest_scoring_entry)
        print("No entries above threshold, appending highest scoring entry: " + highest_scoring_entry["title"])

    relevant_projects = []
    for project_entry in projects:
        if project_entry["title"] in entries_dict:
            score = entries_dict[project_entry["title"]]
            if score > 0.5:
                relevant_projects.append(project_entry)
                print(f"Project: {project_entry['title']}, Score: {score}")

    if len(relevant_projects) == 0:
        # Append the highest scoring project entry if no entries are above the threshold
        highest_scoring_entry = max(projects, key=lambda x: entries_dict.get(x["title"], 0))
        relevant_projects.append(highest_scoring_entry)
        print("No entries above threshold, appending highest scoring entry: " + highest_scoring_entry["title"])

    relevant_leadership = []
    for leadership_entry in leadership:
        if leadership_entry["title"] in entries_dict:
            score = entries_dict[leadership_entry["title"]]
            if score > 0.5:
                relevant_leadership.append(leadership_entry)
                print(f"Leadership: {leadership_entry['title']}, Score: {score}")

    if len(relevant_leadership) == 0:
        # Append the highest scoring leadership entry if no entries are above the threshold
        highest_scoring_entry = max(leadership, key=lambda x: entries_dict.get(x["title"], 0))
        relevant_leadership.append(highest_scoring_entry)
        print("No entries above threshold, appending highest scoring entry: " + highest_scoring_entry["title"])

    relevant_awards = []
    for award_entry in awards:
        if award_entry["title"] in entries_dict:
            score = entries_dict[award_entry["title"]]
            if score > 0.5:
                relevant_awards.append(award_entry)
                print(f"Award: {award_entry['title']}, Score: {score}")

    if len(relevant_awards) == 0:
        # Append the highest scoring award entry if no entries are above the threshold
        highest_scoring_entry = max(awards, key=lambda x: entries_dict.get(x["title"], 0))
        relevant_awards.append(highest_scoring_entry)
        print("No entries above threshold, appending highest scoring entry: " + highest_scoring_entry["title"])

    flat_skills = [s for category in skills_list.values() for s in category]
    relevant_skills = []
    for skill in flat_skills:
        if skill in skills_dict:
            score = skills_dict[skill]
            if score > 0.5:
                relevant_skills.append(skill)
                print(f"Skill: {skill}, Score: {score}")

    if len(relevant_skills) == 0:
        highest_scoring_skill = max(flat_skills, key=lambda x: skills_dict.get(x, 0))
        relevant_skills.append(highest_scoring_skill)
        print("No skills above threshold, appending highest scoring skill: " + highest_scoring_skill)

    # --- Build DOCX ---
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    ACCENT = RGBColor(0xD9, 0x2B, 0x3A)

    for style_name in ("Normal", "List Bullet"):
        fmt = doc.styles[style_name].paragraph_format
        fmt.space_before = Pt(0)
        fmt.space_after = Pt(0)

    # Name header
    name_para = doc.add_paragraph()
    name_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_run = name_para.add_run(name)
    name_run.bold = True
    name_run.font.size = Pt(20)
    name_run.font.color.rgb = ACCENT

    # Contact line
    contact_para = doc.add_paragraph()
    contact_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact_para.add_run(f"{location}  |  {email}")

    def add_section_heading(doc, text):
        p = doc.add_paragraph()
        run = p.add_run(text.upper())
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = ACCENT
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.keep_with_next = True
        return p

    def add_bullet(doc, text):
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(text)
        return p

    # Education
    add_section_heading(doc, "Education")
    edu_para = doc.add_paragraph()
    edu_run = edu_para.add_run(f"{education['institution']}  —  {education['degree']}")
    edu_run.bold = True
    doc.add_paragraph(f"GPA: {education['gpa']}  |  Graduated: {education['graduated']}")
    if education.get("coursework"):
        doc.add_paragraph("Relevant Coursework: " + ", ".join(education["coursework"]))

    def chain_entry(paragraphs):
        for p in paragraphs[:-1]:
            p.paragraph_format.keep_with_next = True

    # Experience
    add_section_heading(doc, "Experience")
    for entry in relevant_experience:
        paras = []
        p = doc.add_paragraph()
        title_run = p.add_run(entry["title"])
        title_run.bold = True
        p.add_run(f"  —  {entry['company']}  ({entry['start']} – {entry['end']})")
        paras.append(p)
        for bullet in entry.get("bullets", []):
            paras.append(add_bullet(doc, bullet))
        chain_entry(paras)

    # Projects
    add_section_heading(doc, "Projects")
    for entry in relevant_projects:
        paras = []
        p = doc.add_paragraph()
        proj_run = p.add_run(entry["title"])
        proj_run.bold = True
        if entry.get("tags"):
            p.add_run("  |  " + ", ".join(entry["tags"]))
        paras.append(p)
        if entry.get("description"):
            paras.append(doc.add_paragraph(entry["description"]))
        chain_entry(paras)

    # Leadership
    if relevant_leadership:
        add_section_heading(doc, "Leadership")
        for entry in relevant_leadership:
            paras = []
            p = doc.add_paragraph()
            title_run = p.add_run(entry["title"])
            title_run.bold = True
            if entry.get("start"):
                p.add_run(f"  ({entry['start']} – {entry['end']})")
            paras.append(p)
            for bullet in entry.get("bullets", []):
                paras.append(add_bullet(doc, bullet))
            chain_entry(paras)

    # Awards
    if relevant_awards:
        add_section_heading(doc, "Awards & Honors")
        for entry in relevant_awards:
            paras = []
            p = doc.add_paragraph()
            award_run = p.add_run(entry["title"])
            award_run.bold = True
            if entry.get("organization"):
                p.add_run(f"  —  {entry['organization']}")
            paras.append(p)
            if entry.get("description"):
                paras.append(doc.add_paragraph(entry["description"]))
            chain_entry(paras)

    # Skills
    add_section_heading(doc, "Skills")
    doc.add_paragraph(", ".join(relevant_skills))

    out_path = f"{name.replace(' ', '_')}_Resume_{datetime.now().strftime('%Y-%m-%d')}.docx"
    try:
        doc.save(out_path)
    except PermissionError:
        out_path = f"{name.replace(' ', '_')}_Resume_{datetime.now().strftime('%Y-%m-%d_%H-%M-%S')}.docx"
        doc.save(out_path)
    print(f"\nSaved: {out_path}")
